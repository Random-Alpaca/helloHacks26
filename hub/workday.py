import openpyxl

from hub.logic import normalise_course_code
from hub.models import Course

_SHEET_NAME = "View My Courses"
_COURSE_HEADER_ALIASES = {"course listing", "course"}


def _find_header_row(rows):
    for row_index, row in enumerate(rows):
        for col_index, cell in enumerate(row):
            value = str(cell).strip().lower() if cell is not None else ""
            if value in _COURSE_HEADER_ALIASES:
                return row_index, col_index
    return None, None


def _course_from_listing(listing, term):
    code_part, _, title = listing.partition(" - ")
    faculty, number, section = normalise_course_code(code_part)
    if not faculty or not number:
        return None

    return Course(
        key=f"UBCV,{term},{faculty},{faculty}{number},{section}",
        code=f"{faculty} {number}",
        section=section,
        term=term,
        title=title.strip() or listing,
        sources=["workday"],
    )


def parse_workday_courses(path, term):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[_SHEET_NAME] if _SHEET_NAME in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    header_row_index, course_col = _find_header_row(rows)
    if header_row_index is None:
        return []

    courses = []
    for row in rows[header_row_index + 1:]:
        if course_col >= len(row) or not row[course_col]:
            continue
        course = _course_from_listing(str(row[course_col]).strip(), term)
        if course is not None:
            courses.append(course)
    return courses
